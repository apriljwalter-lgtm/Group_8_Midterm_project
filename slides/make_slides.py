"""Builds figures and Slides.pptx for the Group 8 presentation. Run from the slides/ folder."""
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

NAMES = "Ryan Merlino, April Walter, Nikolas Platon"
NAVY = RGBColor(0x12, 0x2B, 0x55)
df = pd.read_csv("../diabetes_130.csv")
df["flag"] = (df["readmitted"] != "Not readmitted").astype(int)
df["dx_group"] = pd.cut(df["number_diagnoses"], [0, 4, 6, 8, 99], labels=["1-4", "5-6", "7-8", "9+"])
adm = ["Elective", "Urgent", "Emergency"]
out = ["Not readmitted", "After 30 days", "Within 30 days"]

def save(name):
    plt.tight_layout(); plt.savefig(name, dpi=200); plt.close()

plt.figure(figsize=(9, 5))
sns.histplot(data=df, x="number_diagnoses", discrete=True, color="#2a6f97")
plt.title("Distribution of Number of Diagnoses per Encounter")
plt.xlabel("Number of diagnoses"); plt.ylabel("Number of patients"); save("fig1.png")

plt.figure(figsize=(9, 5))
sns.boxplot(data=df, x="readmitted", y="number_diagnoses", order=out, palette=["#8fb8d8", "#2a6f97", "#d1495b"], hue="readmitted", legend=False)
plt.title("Number of Diagnoses by Readmission Outcome")
plt.xlabel("Readmission outcome"); plt.ylabel("Number of diagnoses"); save("fig2.png")

plt.figure(figsize=(9, 5))
sns.boxplot(data=df, x="admission_type", y="number_diagnoses", hue="readmitted", order=adm, hue_order=out)
plt.title("Number of Diagnoses by Admission Type and Readmission Outcome")
plt.xlabel("Admission type"); plt.ylabel("Number of diagnoses")
plt.legend(title="Readmission outcome", loc="upper left", bbox_to_anchor=(1, 1)); save("fig5.png")

r = df.groupby("number_diagnoses")["flag"].mean().mul(100).reset_index()
r = r[r["number_diagnoses"].between(2, 9)]
plt.figure(figsize=(9, 5))
sns.scatterplot(data=r, x="number_diagnoses", y="flag", s=90, color="#2a6f97")
plt.title("Readmission Rate by Number of Diagnoses")
plt.xlabel("Number of diagnoses"); plt.ylabel("Readmission rate (%)"); plt.ylim(0, 60); save("fig3.png")

g = df.groupby(["admission_type", "dx_group"], observed=True)["flag"].mean().mul(100).reset_index()
plt.figure(figsize=(9, 5))
sns.lineplot(data=g, x="dx_group", y="flag", hue="admission_type", hue_order=adm, marker="o")
plt.title("Readmission Rate by Number of Diagnoses and Admission Type")
plt.xlabel("Number of diagnoses (grouped)"); plt.ylabel("Readmission rate (%)"); plt.ylim(0, 60)
plt.legend(title="Admission type"); save("fig4.png")

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
TEAL = RGBColor(0x5F, 0xCF, 0xC4)
GRAY = RGBColor(0x33, 0x33, 0x33)

def box(s, x, y, w, h, fill):
    r = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background()
    return r

def text(s, t, x, y, w, h, size=24, bold=False, color=GRAY):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.size = Pt(size); p.font.bold = bold; p.space_after = Pt(14)
        p.font.color.rgb = color

def slide(title):
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, 13.333, 1.15, NAVY); box(s, 0, 1.15, 13.333, 0.07, TEAL)
    text(s, title, 0.6, 0.18, 12.2, 0.9, 34, True, RGBColor(255, 255, 255))
    text(s, NAMES + "   |   CAI5735 Mid-Term   |   Group 8", 0.6, 7.0, 12, 0.4, 12, False, RGBColor(0x77, 0x77, 0x77))
    return s

def bullets(s, items, x=0.8, y=1.7, w=11.8, size=26):
    text(s, items, x, y, w, 5, size)

def fig(title, img, items):
    s = slide(title)
    s.shapes.add_picture(img, Inches(0.4), Inches(1.5), width=Inches(8.4))
    text(s, items, 9.0, 1.9, 4.0, 5, 22)

s = prs.slides.add_slide(blank)
box(s, 0, 0, 13.333, 7.5, NAVY); box(s, 0, 4.6, 13.333, 0.08, TEAL)
text(s, "Diagnoses, Admission Type, and Hospital Readmission", 0.8, 1.6, 11.8, 2.5, 46, True, RGBColor(255, 255, 255))
text(s, [NAMES, "CAI5735: Applied Data Science in Health  |  Group 8"], 0.8, 4.9, 11.8, 2, 24, False, RGBColor(255, 255, 255))

s = slide("Research Question")
bullets(s, ["How is the number of diagnoses related to hospital readmission, and does that relationship differ by admission type (Emergency, Urgent, Elective)?",
            "Columns: number_diagnoses, readmitted, admission_type"], y=1.9, size=30)

s = slide("Dataset")
bullets(s, ["3,000 hospital stays of patients with diabetes, from 130 US hospitals (course sample, no missing values)",
            "Source: UCI Machine Learning Repository; Strack et al. (2014), doi:10.24432/C5230J",
            "number_diagnoses: diagnoses recorded for the stay",
            "readmitted: not readmitted (60%), after 30 days (31%), within 30 days (9%)",
            "admission_type: Emergency (57%), Elective (22%), Urgent (21%)"], size=24)

s = slide("Analysis Process")
bullets(s, ["1. Kept the three relevant columns",
            "2. Added computed columns: readmitted flag (0/1) and diagnosis group (1-4, 5-6, 7-8, 9+)",
            "3. Filtered to Emergency admissions as a check within one type",
            "4. Grouped by admission type and diagnosis group; sorted by readmission rate",
            "5. Made a histogram, boxplot, scatter plot, and group comparison charts"], size=26)

fig("Most stays have 9 diagnoses", "fig1.png", ["1,376 of 3,000 stays have exactly 9 diagnoses", "Likely because the source hospital systems capped recorded diagnoses, so patients with more conditions may be undercounted"])
fig("Readmitted patients have more diagnoses", "fig2.png", ["Average diagnoses: 7.0 for not readmitted vs 7.7 for readmitted (either timing)"])
fig("Readmitted patients: more diagnoses in all types", "fig5.png", ["Average diagnoses, not readmitted vs readmitted (either timing):", "Elective: 6.7 vs 7.4-7.7", "Urgent: 6.8 vs 7.4-7.7", "Emergency: 7.2 vs 7.7", "The gap is smallest for Emergency"])
fig("Readmission rises with diagnoses", "fig3.png", ["Rate climbs from 30% at 5 diagnoses to 46% at 9", "Points at 2-3 diagnoses are noisy (39 and 102 patients)", "Overall: 40.0% readmitted"])
fig("Same upward trend in each admission type", "fig4.png", ["1-4 diagnoses: 24.6% readmitted; 9+: 46.0%", "Elective is lowest in every group (34.1% vs 41.4% Urgent, 41.8% Emergency)", "Emergency-only check shows the same rise"])

s = slide("Key Takeaways")
bullets(s, ["More diagnoses goes with higher readmission, leveling off after about 7",
            "Readmission rises in all three admission types; Elective is consistently lowest",
            "Limits: association, not causation; many stays have exactly 9 diagnoses; small subgroups (Urgent 1-4 has 74 patients)"], size=28)

s = prs.slides.add_slide(blank)
box(s, 0, 0, 13.333, 7.5, NAVY)
text(s, "Questions?", 0.8, 2.4, 11.8, 1.5, 54, True, RGBColor(255, 255, 255))
text(s, "github.com/apriljwalter-lgtm/Group_8_Midterm_project", 0.8, 4.2, 11.8, 1, 24, False, RGBColor(255, 255, 255))
prs.save("Slides.pptx")
